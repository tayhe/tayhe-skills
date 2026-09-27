# -*- coding: utf-8 -*-
"""
「剧场与纸」讲义 PPT 模板（lecture-ppt-builder skill 资产）
============================================================
用法：把本文件整体复制为项目根的 build_<讲名>.py，然后：
  1. 修改「本讲配置」区（OUT、FOOTER、封面 kicker）；
  2. 在 IMG 字典注册本讲要用的图片（键名自取，值为 Path）；
  3. 新增图片必须在 FOCUS 表登记焦点 (fx, fy)，并用渲染预览验证面部；
  4. 在文件末尾「幻灯片内容」区按手稿标题层级逐页编写。

铁律（详见 skill 的 references/design-system.md）：
  · 一页一个论点；引文必带行号出处；
  · 页眉 kicker 一律带手稿编号（1.1 / 2.3 / 3.4·a / 3.4 小结…），
    页面顺序与手稿标题层级逐节对应，手稿的 a–d 枚举与「小结」必须显式保留；
  · 正文不得低于 15pt，引文不得低于 20pt（教室投影硬指标）；
  · 只用 assets/hd_images/ 的考据图，禁止用 assets/bg_images/ 低清碎片。
"""
from pathlib import Path
import tempfile
from PIL import Image

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 本讲配置
HERE = Path(__file__).resolve().parent
HD = HERE / "assets" / "hd_images"
OUT = HERE / "出讲义" / "第X讲_标题.pptx"          # TODO 改成本讲文件名
TMP = Path(tempfile.mkdtemp(prefix="lec_crops_"))

FOOTER = "第X讲 · 副标题"                          # TODO 页脚文字
COVER_KICKER = "政治哲学 · 第X讲"                  # TODO 封面眉标

# 素材分组（assets/hd_images/ 下，按主题分组的考据高清图）
PHIL    = HD / "01_哲学思想与范式"
CIV     = HD / "02_希腊地理与文明演进"
CITY    = HD / "03_城邦制度与公民实践"
TRAGEDY = HD / "04_希腊悲剧与法治创制"
ATREUS  = HD / "神话专题_阿特柔斯家族"
THEBES  = HD / "神话专题_忒拜家族"
GODS    = HD / "神话专题_诸神_特洛伊与忒拜"


def precrop(src, left, top, right, bottom):
    """素材照片带画框/白边时先按比例预裁，返回临时文件路径。
    例：LYTRAS = precrop(THEBES / "16_...Lytras.jpg", 0.07, 0.065, 0.93, 0.94)"""
    im = Image.open(src)
    w, h = im.size
    out = TMP / (src.stem + "_crop.jpg")
    im.crop((int(w * left), int(h * top), int(w * right), int(h * bottom))
            ).convert("RGB").save(out, "JPEG", quality=92)
    return out


IMG = {
    # TODO 注册本讲图片，例：
    # "lytras": precrop(THEBES / "16_忒拜_利特拉斯_安提戈涅在波吕尼刻斯遗体前_Lytras.jpg",
    #                   0.07, 0.065, 0.93, 0.94),   # 原图带金色画框，必须预裁
    # "moreau": THEBES / "13_忒拜_莫罗_俄狄浦斯与斯芬克斯_Moreau.jpg",
}

# ---------------------------------------------------------------- 色彩
PAPER    = RGBColor(0xF7, 0xF2, 0xE8)   # 暖象牙纸底
PAPER_D  = RGBColor(0xEF, 0xE7, 0xD6)
INK      = RGBColor(0x28, 0x24, 0x1E)   # 墨黑
INK_SOFT = RGBColor(0x54, 0x4D, 0x43)
GRAY     = RGBColor(0x8B, 0x82, 0x74)
TERRA    = RGBColor(0xA5, 0x3C, 0x24)   # 陶瓶红
TERRA_D  = RGBColor(0x7C, 0x2C, 0x19)
GOLD     = RGBColor(0xA8, 0x82, 0x2E)   # 青铜金
NIGHT    = RGBColor(0x1E, 0x1A, 0x16)   # 暖黑（深色页/遮罩）
NIGHT_2  = RGBColor(0x2A, 0x24, 0x1E)
LINE     = RGBColor(0xD9, 0xCD, 0xB6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CREAM    = RGBColor(0xF3, 0xEA, 0xD9)

# ---------------------------------------------------------------- 字体
F_SERIF = "Noto Serif SC"    # 思源宋体：标题/正文
F_SANS  = "Noto Sans SC"     # 思源黑体：眉标/注释
F_KAI   = "华文楷体"          # 引文
F_LATIN = "Times New Roman"  # 西文/数字

# ---------------------------------------------------------------- 画布
W, H = 13.333, 7.5   # 16:9，英寸
M = 0.62             # 外边距
FOOT_Y = 7.12

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)


# ================================================================ 基础工具
def _cjk(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    succ = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *succ[tag])
        el.set("typeface", name)


def _spc(run, val):
    run._r.get_or_add_rPr().set("spc", str(val))


def rect(slide, x, y, w, h, color, line_color=None, line_w=None, shadow=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w or 0.75)
    sp.shadow.inherit = False
    return sp


def alpha_rect(slide, x, y, w, h, color, alpha):
    """半透明遮罩。alpha 0~100（数值越大越不透明）"""
    sp = rect(slide, x, y, w, h, color)
    sf = sp.fill._xPr.find(qn("a:solidFill"))
    clr = sf.find(qn("a:srgbClr"))
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))})
    clr.append(a)
    return sp


def hline(slide, x, y, w, color, weight=1.0):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln


def txt(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """通用文本框。paras 为 dict 列表：
    dict(text=..., font=, size=, color=, bold=, italic=, align=, spacing=, before=, after=, spc=)
    或 dict(runs=[{text, font, size, color, bold, italic, spc}, ...], ...) 混排"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("spacing"):
            para.line_spacing = p["spacing"]
        para.space_before = Pt(p.get("before", 0))
        para.space_after = Pt(p.get("after", 0))
        runs = p.get("runs") or [{"text": p.get("text", ""), **{k: p[k] for k in
             ("font", "size", "color", "bold", "italic", "spc") if k in p}}]
        for rs in runs:
            r = para.add_run()
            r.text = rs["text"]
            f = r.font
            f.size = Pt(rs.get("size", p.get("size", 14)))
            f.bold = rs.get("bold", p.get("bold", False))
            f.italic = rs.get("italic", p.get("italic", False))
            f.color.rgb = rs.get("color", p.get("color", INK))
            _cjk(r, rs.get("font", p.get("font", F_SERIF)))
            spc = rs.get("spc", p.get("spc"))
            if spc:
                _spc(r, spc)
    return tb


# ---------------------------------------------------------------- 焦点表
# 各图焦点 (fx, fy, 0~1)：裁切窗口以此为中心，保住人物面部/画面重点。
# 默认 (0.5, 0.38)——古典油画/立像面部多在画面上 1/3 处。
# 新增图片必须登记并渲染验证面部；预裁后的临时图按 <原名>_crop.jpg 命名登记。
FOCUS = {
    # —— 阿特柔斯家族（第一讲 A 用）
    "01_阿特柔斯_始祖之罪_坦塔罗斯冥界受难_Assereto.jpg":                 (0.50, 0.30),
    "02_阿特柔斯_提埃波罗_献祭伊菲革涅亚_Tiepolo.jpg":                    (0.50, 0.46),
    "03_阿特柔斯_桑兹_绝望的预言者卡珊德拉_Sandys.jpg":                   (0.50, 0.28),
    "04_阿特柔斯_约翰科利尔_行刺后的克吕泰涅斯特拉_Collier.jpg":          (0.50, 0.26),
    "05_阿特柔斯_盖兰_克吕泰涅斯特拉行刺沉睡的阿伽门农_Guerin.jpg":       (0.45, 0.32),
    "06_阿特柔斯_莱顿_厄勒克特拉在阿伽门农墓前奠酒_Leighton.jpg":         (0.50, 0.34),
    "07_阿特柔斯_布格罗_复仇女神追逐俄瑞斯忒斯_Bouguereau.jpg":           (0.50, 0.40),
    "08_阿特柔斯_萨金特_被复仇女神追逐的俄瑞斯忒斯_Sargent.jpg":          (0.50, 0.38),
    "09_阿特柔斯_古希腊红绘双耳大罐_俄瑞斯忒斯避难特尔斐阿波罗圣所.jpg":  (0.50, 0.50),
    "10_阿特柔斯_雅典战神山全景_雅典娜创立阿瑞奥帕戈斯法庭.jpg":          (0.50, 0.50),
    "25_布格罗_复仇女神追逐俄瑞斯忒斯_Bouguereau_Orestes.jpg":            (0.50, 0.40),
    # —— 诸神·特洛伊与忒拜
    "02_宙斯与忒提斯_安格尔_格拉内博物馆.jpg":                            (0.50, 0.30),
    "04_命运天平_弗斯利_灵魂与命运的秤量_苏黎世美术馆.jpg":               (0.50, 0.35),
    "05_赫拉_坎帕纳大理石巨幅胸像_卢浮宫.jpg":                            (0.50, 0.32),
    "07_雅典娜_吉斯蒂尼亚尼大理石立像_梵蒂冈博物馆.jpg":                  (0.50, 0.28),
    "08_雅典娜阻止阿喀琉斯拔剑刺杀阿伽门农_庞贝壁画_那不勒斯考古博物馆.jpg": (0.42, 0.18),
    "10_阿波罗_贝尔维德尔大理石立像_梵蒂冈博物馆.jpg":                    (0.50, 0.28),
    "11_阿波罗德尔斐女祭司皮提亚_约翰柯利尔_南澳美术馆.jpg":              (0.45, 0.20),
    "12_阿佛洛狄忒_米洛的维纳斯_卢浮宫.jpg":                              (0.50, 0.28),
    "13_帕里斯的评判_鲁本斯_英国国家美术馆.jpg":                          (0.50, 0.38),
    "14_阿瑞斯_卢多维西大理石雕像_罗马国家博物馆.jpg":                    (0.45, 0.32),
    "18_波塞冬_阿特米西昂青铜神像_雅典国家考古博物馆.jpg":                (0.50, 0.28),
    "22_阿尔忒弥斯_凡尔赛的狄安娜_卢浮宫.jpg":                            (0.50, 0.28),
    "23_戴安娜与阿克泰翁_提香_英国国家美术馆.jpg":                        (0.50, 0.38),
    "29_普路托掠夺普罗瑟平娜_贝尔尼尼_博尔盖塞美术馆.jpg":                (0.50, 0.30),
    "32_阿喀琉斯战车拖拽赫克托尔遗体_马奇_阿喀琉斯宫壁画.jpg":             (0.50, 0.50),
    "37_雅典娜_韦莱特里的帕拉斯大理石立像_卢浮宫.jpg":                    (0.50, 0.28),
    "45_赫拉哺乳赫拉克勒斯与银河起源_丁托列托_英国国家美术馆.jpg":        (0.50, 0.30),
    "51_俄狄浦斯破解斯芬克斯之谜_古希腊红绘基里克斯陶杯_梵蒂冈博物馆.jpg":  (0.50, 0.45),
    "52_俄狄浦斯与安提戈涅流亡_埃克斯伯格_丹麦国立美术馆.jpg":              (0.50, 0.40),
    # —— 忒拜家族（第一讲 B 用）
    "12_忒拜_安格尔_俄狄浦斯与斯芬克斯之谜_Ingres.jpg":                   (0.40, 0.32),
    "13_忒拜_莫罗_俄狄浦斯与斯芬克斯_Moreau.jpg":                         (0.50, 0.20),
    "14_忒拜_吉鲁斯特_盲目的俄狄浦斯在科罗诺斯_Giroust.jpg":               (0.55, 0.45),
    "15_忒拜_特申多夫_俄狄浦斯王与陪伴他的安提戈涅_Teschendorff.jpg":      (0.50, 0.22),
    "16_忒拜_利特拉斯_安提戈涅在波吕尼刻斯遗体前_Lytras.jpg":              (0.45, 0.42),
    "lytras_noframe.jpg":                                                 (0.45, 0.42),
    "17_忒拜_贡斯当_安提戈涅守护波吕尼刻斯_Benjamin_Constant.jpg":         (0.45, 0.40),
    "18_忒拜_阿比尔高_伊斯墨涅与安提戈涅向忒修斯求助_Abildgaard.jpg":      (0.50, 0.50),
    # —— 悲剧与法治 / 城邦 / 文明 / 哲学
    "13_阿伽门农黄金面具_Mask_of_Agamemnon.jpg":                          (0.50, 0.45),
    "22_雅典卫城南坡狄奥尼索斯剧场_Theatre_of_Dionysus.jpg":               (0.50, 0.50),
    "23_埃斯库罗斯大理石胸像_Aeschylus_Bust.jpg":                         (0.50, 0.35),
    "15_冯克兰茨_雅典卫城与战神山理想复原图_Acropolis_Klenze.jpg":         (0.50, 0.40),
    "16_普尼克斯山公民大会会址与卫城全景_Pnyx_Acropolis.jpg":              (0.50, 0.50),
    "18_伯利克里大理石胸像_Pericles_Bust.jpg":                            (0.50, 0.32),
    "19_冯福尔茨_伯利克里国葬演说_Pericles_Funeral_Oration.jpg":           (0.58, 0.40),
    "09_爱琴海全景_圣托里尼火山口_Santorini_Caldera.jpg":                 (0.50, 0.45),
    "06_黑格尔油画肖像_Hegel_Schlesinger.jpg":                            (0.45, 0.38),
}


def cover_pic(slide, path, x, y, w, h, focus=None):
    """焦点裁切填充（不拉伸）：裁切窗口以焦点为中心，防止切掉人物面部"""
    im = Image.open(path)
    iw, ih = im.size
    target = w / h
    if focus is None:
        focus = FOCUS.get(Path(path).name, (0.5, 0.38))
    fx, fy = focus
    if iw / ih > target:
        nw = int(ih * target)
        cx = min(max(fx * iw - nw / 2, 0), iw - nw)
        box = (int(cx), 0, int(cx) + nw, ih)
    else:
        nh = int(iw / target)
        cy = min(max(fy * ih - nh / 2, 0), ih - nh)
        box = (0, int(cy), iw, int(cy) + nh)
    out = TMP / f"c{len(list(TMP.iterdir()))}.jpg"
    im.crop(box).convert("RGB").save(out, "JPEG", quality=90)
    return slide.shapes.add_picture(str(out), Inches(x), Inches(y), Inches(w), Inches(h))


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def paper_bg(slide, color=PAPER):
    rect(slide, 0, 0, W, H, color)


def footer(slide, page):
    txt(slide, M, FOOT_Y, 6.0, 0.3, [dict(text=FOOTER, font=F_SANS, size=9.5, color=GRAY, spc=100)])
    txt(slide, W - M - 1.0, FOOT_Y, 1.0, 0.3,
        [dict(text=f"{page:02d}", font=F_LATIN, size=10, color=GRAY, align=PP_ALIGN.RIGHT)])


def header(slide, kicker, title, page, title_size=30):
    """象牙页标准页眉：kicker 必须带手稿编号"""
    txt(slide, M, 0.44, W - 2 * M, 0.32,
        [dict(text=kicker, font=F_SANS, size=12.5, color=TERRA, bold=True, spc=300)])
    txt(slide, M, 0.78, W - 2 * M, 0.68,
        [dict(text=title, font=F_SERIF, size=title_size, color=INK, bold=True)])
    hline(slide, M, 1.58, 0.85, GOLD, 2.2)
    hline(slide, M + 0.85, 1.58, W - 2 * M - 0.85, LINE, 0.9)
    footer(slide, page)


def lead_par(slide, text, y=1.82, x=M, w=None, color=INK_SOFT, size=16):
    txt(slide, x, y, w or (W - 2 * M), 0.55,
        [dict(text=text, font=F_SERIF, size=size, color=color, italic=True, spacing=1.25)])


def bullets(slide, x, y, w, items, size=17, gap=10, spacing=1.32, marker_color=TERRA):
    """要点列表。items: 字符串或 (加粗引导语, 正文) 元组"""
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            runs = [{"text": "▪ ", "font": F_SANS, "size": size - 1.5, "color": marker_color}]
            if head:
                runs.append({"text": head, "size": size, "color": INK, "bold": True, "font": F_SERIF})
            runs.append({"text": body, "size": size, "color": INK, "font": F_SERIF})
            paras.append(dict(runs=runs, spacing=spacing, after=gap))
        else:
            paras.append(dict(runs=[
                {"text": "▪ ", "font": F_SANS, "size": size - 1.5, "color": marker_color},
                {"text": it, "size": size, "color": INK, "font": F_SERIF},
            ], spacing=spacing, after=gap))
    txt(slide, x, y, w, 4.5, paras)


def point_bar(slide, x, y, w, text, size=16.5):
    """页底陶红结论条：每页论点的收束句"""
    rect(slide, x, y, 0.1, 0.72, TERRA)
    txt(slide, x + 0.28, y - 0.02, w - 0.28, 1.0,
        [dict(text=text, font=F_SERIF, size=size, color=TERRA_D, bold=True, spacing=1.3)])


# ================================================================ 页面模板
def s_cover(title, sub, meta, img, en=""):
    s = blank()
    cover_pic(s, img, 0, 0, W, H)
    alpha_rect(s, 0, 0, W, H, NIGHT, 58)
    alpha_rect(s, 0, H - 2.4, W, 2.4, NIGHT, 45)
    txt(s, 0.9, 1.45, 11.5, 0.4,
        [dict(text=COVER_KICKER, font=F_SANS, size=15, color=GOLD, bold=True, spc=500)])
    txt(s, 0.88, 1.95, 11.6, 1.5,
        [dict(text=title, font=F_SERIF, size=46, color=WHITE, bold=True, spc=200)])
    hline(s, 0.92, 3.42, 1.6, GOLD, 2.5)
    txt(s, 0.9, 3.66, 11.5, 0.55,
        [dict(text=sub, font=F_SERIF, size=20, color=CREAM, spacing=1.3)])
    if en:
        txt(s, 0.9, 4.3, 11.5, 0.4,
            [dict(text=en, font=F_LATIN, size=14, color=RGBColor(0xC9, 0xBE, 0xA9), italic=True)])
    txt(s, 0.9, 6.72, 11.5, 0.35,
        [dict(text=meta, font=F_SANS, size=12, color=RGBColor(0xCF, 0xC5, 0xB2), spc=200)])
    return s


def s_section(num, title, sub, img, note=""):
    """章节页：满版油画 + 大序号"""
    s = blank()
    cover_pic(s, img, 0, 0, W, H)
    alpha_rect(s, 0, 0, W, H, NIGHT, 62)
    txt(s, 0.95, 1.55, 3.2, 2.2,
        [dict(text=num, font=F_LATIN, size=96, color=GOLD, bold=True)])
    hline(s, 1.02, 3.8, 1.1, GOLD, 2.2)
    txt(s, 0.98, 4.0, 11.0, 0.95,
        [dict(text=title, font=F_SERIF, size=36, color=WHITE, bold=True, spc=150)])
    txt(s, 1.0, 4.98, 10.8, 0.6,
        [dict(text=sub, font=F_SERIF, size=18, color=CREAM, spacing=1.35)])
    if note:
        txt(s, 1.0, 6.6, 10.8, 0.4,
            [dict(text=note, font=F_SANS, size=12, color=RGBColor(0xB9, 0xAE, 0x9B), spc=150)])
    return s


def s_quote(kicker, quote_paras, source, img, note=""):
    """引文页：满版油画 + 暗化遮罩 + 楷体引文。>=7 行自动缩为 20pt 防溢出。"""
    s = blank()
    cover_pic(s, img, 0, 0, W, H)
    alpha_rect(s, 0, 0, W, H, NIGHT, 70)
    alpha_rect(s, 0, H - 1.7, W, 1.7, NIGHT, 48)
    txt(s, 0.95, 0.72, 11.4, 0.38,
        [dict(text=kicker, font=F_SANS, size=13, color=GOLD, bold=True, spc=350)])
    hline(s, 0.98, 1.2, 0.85, GOLD, 2.0)
    n = len(quote_paras)
    if n <= 6:
        qsize, qsp, qafter, qy = 23, 1.5, 8, 1.7
    else:
        qsize, qsp, qafter, qy = 20, 1.45, 4, 1.5
    paras = [dict(text=q, font=F_KAI, size=qsize, color=WHITE, spacing=qsp, after=qafter)
             for q in quote_paras]
    txt(s, 0.98, qy, 10.2, 4.5, paras)
    txt(s, 0.98, 6.18, 11.0, 0.4,
        [dict(text=source, font=F_SERIF, size=15, color=GOLD)])
    if note:
        txt(s, 0.98, 6.62, 11.4, 0.6,
            [dict(text=note, font=F_SERIF, size=15, color=CREAM, spacing=1.3)])
    return s


def s_dark(kicker, title, body_paras, page, size=18):
    """深色思辨页（纯版式，无图）：核心命题/设问统摄页"""
    s = blank()
    rect(s, 0, 0, W, H, NIGHT)
    rect(s, 0, 0, W, 0.14, TERRA)
    txt(s, M + 0.3, 0.8, W - 2 * M, 0.38,
        [dict(text=kicker, font=F_SANS, size=13, color=GOLD, bold=True, spc=350)])
    txt(s, M + 0.3, 1.28, W - 2 * M - 0.6, 0.8,
        [dict(text=title, font=F_SERIF, size=29, color=WHITE, bold=True, spacing=1.15)])
    hline(s, M + 0.33, 2.28, 0.85, TERRA, 2.4)
    paras = []
    for b in body_paras:
        if isinstance(b, tuple):
            head, body = b
            paras.append(dict(runs=[
                {"text": "▪ ", "font": F_SANS, "size": size - 2, "color": TERRA},
                {"text": head, "size": size, "color": GOLD, "bold": True, "font": F_SERIF},
                {"text": body, "size": size, "color": CREAM, "font": F_SERIF},
            ], spacing=1.45, after=14))
        else:
            paras.append(dict(runs=[
                {"text": "▪ ", "font": F_SANS, "size": size - 2, "color": TERRA},
                {"text": b, "size": size, "color": CREAM, "font": F_SERIF},
            ], spacing=1.45, after=14))
    txt(s, M + 0.3, 2.62, W - 2 * M - 0.6, 4.2, paras)
    txt(s, W - M - 1.0, FOOT_Y, 1.0, 0.3,
        [dict(text=f"{page:02d}", font=F_LATIN, size=10, color=GRAY, align=PP_ALIGN.RIGHT)])
    return s


def s_content(kicker, title, page, lead="", img=None, img_w=4.6, img_cap="", focus=None):
    """象牙图文页；返回 (slide, 正文可用宽度 tw)"""
    s = blank()
    paper_bg(s)
    header(s, kicker, title, page)
    tw = W - 2 * M
    if img:
        iw = img_w
        ix = W - M - iw
        cover_pic(s, img, ix, 1.78, iw, 4.62, focus=focus)
        if img_cap:
            alpha_rect(s, ix, 1.78 + 4.62 - 0.44, iw, 0.44, NIGHT, 55)
            txt(s, ix + 0.12, 1.78 + 4.62 - 0.40, iw - 0.24, 0.36,
                [dict(text=img_cap, font=F_SANS, size=8.5, color=CREAM)])
        tw = ix - M - 0.5
    if lead:
        lead_par(s, lead, w=tw)
    return s, tw


def s_compare(kicker, title, page, left, right, lead="", bottom="", cols_color=(TERRA, GOLD)):
    """双栏对照页。left/right: dict(head, sub, items)"""
    s = blank()
    paper_bg(s)
    header(s, kicker, title, page)
    if lead:
        lead_par(s, lead)
    top = 2.25 if lead else 1.95
    cw = (W - 2 * M - 0.55) / 2
    for i, col in enumerate((left, right)):
        x = M + i * (cw + 0.55)
        rect(s, x, top, cw, 0.58, cols_color[i])
        txt(s, x + 0.2, top + 0.1, cw - 0.4, 0.42,
            [dict(text=col["head"], font=F_SERIF, size=17, color=WHITE, bold=True)])
        rect(s, x, top + 0.58, cw, 0.4, PAPER_D)
        txt(s, x + 0.2, top + 0.645, cw - 0.4, 0.32,
            [dict(text=col["sub"], font=F_SANS, size=12, color=INK_SOFT, spc=120)])
        rect(s, x, top + 0.98, cw, 2.92, RGBColor(0xFC, 0xFA, 0xF5), line_color=LINE, line_w=0.75)
        paras = []
        for it in col["items"]:
            if isinstance(it, tuple):
                head, body = it
                paras.append(dict(runs=[
                    {"text": head, "size": 16, "color": cols_color[i], "bold": True, "font": F_SERIF},
                    {"text": body, "size": 16, "color": INK, "font": F_SERIF},
                ], spacing=1.28, after=8))
            else:
                paras.append(dict(text=it, font=F_SERIF, size=16, color=INK, spacing=1.28, after=8))
        txt(s, x + 0.24, top + 1.2, cw - 0.48, 2.6, paras)
    if bottom:
        point_bar(s, M, top + 4.12, W - 2 * M, bottom)
    return s


def s_cards(kicker, title, page, cards, lead="", cols=2, bottom=""):
    """卡片页。cards: (序号/标签, 小标题, 正文)。正文较长（>80字/卡）时改用 s_content+bullets。"""
    s = blank()
    paper_bg(s)
    header(s, kicker, title, page)
    if lead:
        lead_par(s, lead)
    top = 2.34 if lead else 2.0
    n = len(cards)
    rows = (n + cols - 1) // cols
    gap = 0.35
    cw = (W - 2 * M - (cols - 1) * gap) / cols
    ch = min(2.3 if rows == 1 else 1.9, (4.6 - (rows - 1) * gap) / rows)
    for i, (tag, head, body) in enumerate(cards):
        r, c = divmod(i, cols)
        x = M + c * (cw + gap)
        y = top + r * (ch + gap)
        rect(s, x, y, cw, ch, RGBColor(0xFC, 0xFA, 0xF5), line_color=LINE, line_w=0.75)
        rect(s, x, y, 0.08, ch, TERRA)
        txt(s, x + 0.24, y + 0.16, cw - 0.46, 0.4, [dict(runs=[
            {"text": tag + "  ", "font": F_LATIN, "size": 14, "color": GOLD, "bold": True},
            {"text": head, "font": F_SERIF, "size": 16, "color": INK, "bold": True},
        ])])
        txt(s, x + 0.24, y + 0.6, cw - 0.46, ch - 0.74,
            [dict(text=body, font=F_SERIF, size=15, color=INK_SOFT, spacing=1.26)])
    if bottom:
        point_bar(s, M, top + rows * ch + (rows - 1) * gap + 0.2, W - 2 * M, bottom)
    return s


# ================================================================ 页码
PAGE = 0
def pg():
    global PAGE
    PAGE += 1
    return PAGE


# ================================================================ 幻灯片内容
# 在此按手稿标题层级逐页编写。每种页面模板对应一种论证功能：
#   s_cover    封面（满版画）
#   s_section  章节页（满版画 + 大序号）
#   s_quote    引文页（满版画 + 楷体，带行号出处）
#   s_dark     深色思辨页（核心命题 / 设问统摄后续分页）
#   s_content  象牙图文页（lead + bullets + point_bar）
#   s_compare  双栏对照页（概念对决）
#   s_cards    卡片页（并列要点，正文勿超 80 字/卡）
#
# 示例：
# pg()
# s_cover("标题", "副标题", "西南石油大学 · 政治哲学课程　何舒骏", IMG["x"], en="...")
# s_dark("导言 · 本讲问题", "设问句", [("设问一：", "……"), ...], pg())
# s, tw = s_content("2.1 · kicker", "页面标题", pg(), lead="……", img=IMG["x"], img_cap="图注")
# bullets(s, M, 2.55, tw, [("引导语：", "正文"), ...], size=16, gap=10)
# point_bar(s, M, 6.0, tw, "结论句")
#
# 注意：本讲所有页面代码必须写在下面 prs.save 之前。

prs.save(str(OUT))
print("saved:", OUT, "slides:", PAGE)
